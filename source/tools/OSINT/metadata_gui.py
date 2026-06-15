import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import os
import sys
import json
import threading
import re
import zipfile
import xml.etree.ElementTree as ET

# Optional dependencies for parsing
try:
    import PyPDF2
    has_pdf = True
except ImportError:
    has_pdf = False

try:
    from docx import Document
    has_docx = True
except ImportError:
    has_docx = False

try:
    import openpyxl
    has_xlsx = True
except ImportError:
    has_xlsx = False

try:
    from pptx import Presentation
    has_pptx = True
except ImportError:
    has_pptx = False

try:
    from PIL import Image
    from PIL.ExifTags import TAGS, GPSTAGS
    has_pil = True
except ImportError:
    has_pil = False

class MetadataScraperGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("SuperSploit Metadata Scraper - DEEP INSPECTION")
        self.root.geometry("1000x800")
        self.root.configure(bg="#f0f0f0")

        self.style = ttk.Style()
        self.style.theme_use('clam')

        self.results = {}
        self.files_to_scan = []

        self._build_ui()

    def _build_ui(self):
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        # --- 1. File Selection ---
        selection_frame = ttk.LabelFrame(main_frame, text=" 📂 File Selection ", padding="10")
        selection_frame.pack(fill=tk.X, pady=5)

        self.path_var = tk.StringVar()
        ttk.Entry(selection_frame, textvariable=self.path_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Button(selection_frame, text="Browse File", command=self._browse_file).pack(side=tk.LEFT, padx=2)
        ttk.Button(selection_frame, text="Browse Folder", command=self._browse_folder).pack(side=tk.LEFT, padx=2)

        # --- 2. Action Buttons ---
        btn_frame = ttk.Frame(main_frame, padding="5")
        btn_frame.pack(fill=tk.X)

        self.analyze_btn = ttk.Button(btn_frame, text="🔍 DEEP ANALYZE", command=self._start_analysis)
        self.analyze_btn.pack(side=tk.LEFT, padx=5)
        
        self.export_btn = ttk.Button(btn_frame, text="📥 EXPORT RESULTS", command=self._export_results, state=tk.DISABLED)
        self.export_btn.pack(side=tk.LEFT, padx=5)

        ttk.Button(btn_frame, text="🧹 CLEAR", command=self._clear_all).pack(side=tk.RIGHT, padx=5)

        # --- 3. Results Area (Paned Window) ---
        paned = ttk.PanedWindow(main_frame, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, pady=5)

        # Left: File List
        list_frame = ttk.Frame(paned, padding="5")
        paned.add(list_frame, weight=1)

        ttk.Label(list_frame, text="Scanned Files:").pack(anchor=tk.W)
        self.file_list = tk.Listbox(list_frame, font=("Courier", 10))
        self.file_list.pack(fill=tk.BOTH, expand=True)
        self.file_list.bind('<<ListboxSelect>>', self._on_file_select)

        # Right: Metadata View
        details_frame = ttk.Frame(paned, padding="5")
        paned.add(details_frame, weight=2)

        ttk.Label(details_frame, text="Deep Analysis Details:").pack(anchor=tk.W)
        self.details_text = scrolledtext.ScrolledText(details_frame, bg="white", font=("Courier", 10))
        self.details_text.pack(fill=tk.BOTH, expand=True)

        # --- 4. Status Bar ---
        self.status_var = tk.StringVar(value="Ready")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)

        # Check for dependencies
        self._check_dependencies()

    def _check_dependencies(self):
        missing = []
        if not has_pdf: missing.append("PyPDF2")
        if not has_docx: missing.append("python-docx")
        if not has_xlsx: missing.append("openpyxl")
        if not has_pptx: missing.append("python-pptx")
        if not has_pil: missing.append("Pillow")
        
        if missing:
            self.details_text.insert(tk.END, "⚠️ WARNING: Missing dependencies for full support:\n")
            for m in missing:
                self.details_text.insert(tk.END, f"  - {m}\n")
            self.details_text.insert(tk.END, "\nInstall them with: pip install " + " ".join(missing) + "\n\n")

    def _browse_file(self):
        filename = filedialog.askopenfilename(filetypes=[
            ("All Supported", "*.pdf *.docx *.xlsx *.pptx *.jpg *.jpeg *.png"),
            ("Documents", "*.pdf *.docx *.xlsx *.pptx"),
            ("Images", "*.jpg *.jpeg *.png"),
            ("All Files", "*.*")
        ])
        if filename:
            self.path_var.set(filename)

    def _browse_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.path_var.set(folder)

    def _clear_all(self):
        self.results = {}
        self.files_to_scan = []
        self.file_list.delete(0, tk.END)
        self.details_text.delete(1.0, tk.END)
        self.export_btn.config(state=tk.DISABLED)
        self.status_var.set("Ready")
        self._check_dependencies()

    def _get_office_xml_props(self, file_path):
        """Peeks inside the ZIP structure of Office docs for hidden XML metadata."""
        props = {}
        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                # Check app.xml for Company, Template, App Version
                if 'docProps/app.xml' in z.namelist():
                    xml = z.read('docProps/app.xml')
                    root = ET.fromstring(xml)
                    for child in root:
                        tag = child.tag.split('}')[-1]
                        if child.text: props[f"Office_{tag}"] = child.text
                
                # Check core.xml for deep timestamps
                if 'docProps/core.xml' in z.namelist():
                    xml = z.read('docProps/core.xml')
                    root = ET.fromstring(xml)
                    for child in root:
                        tag = child.tag.split('}')[-1]
                        if child.text: props[f"Core_{tag}"] = child.text
        except: pass
        return props

    def _scrape_pdf(self, file_path):
        if not has_pdf: return {"error": "PyPDF2 not installed"}
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                info = reader.metadata
                return {str(k): str(v) for k, v in info.items()}
        except Exception as e:
            return {"error": str(e)}

    def _scrape_docx(self, file_path):
        if not has_docx: return {"error": "python-docx not installed"}
        try:
            doc = Document(file_path)
            prop = doc.core_properties
            meta = {
                "Author": prop.author,
                "Created": str(prop.created),
                "Last Modified By": prop.last_modified_by,
                "Revision": prop.revision,
                "Title": prop.title
            }
            meta.update(self._get_office_xml_props(file_path))
            return meta
        except Exception as e:
            return {"error": str(e)}

    def _scrape_xlsx(self, file_path):
        if not has_xlsx: return {"error": "openpyxl not installed"}
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            prop = wb.properties
            meta = {
                "Creator": prop.creator,
                "Last Modified By": prop.lastModifiedBy,
                "Created": str(prop.created),
                "Modified": str(prop.modified),
                "Title": prop.title
            }
            meta.update(self._get_office_xml_props(file_path))
            return meta
        except Exception as e:
            return {"error": str(e)}

    def _scrape_pptx(self, file_path):
        if not has_pptx: return {"error": "python-pptx not installed"}
        try:
            prs = Presentation(file_path)
            prop = prs.core_properties
            meta = {
                "Author": prop.author,
                "Created": str(prop.created),
                "Last Modified By": prop.last_modified_by,
                "Revision": prop.revision,
                "Title": prop.title
            }
            meta.update(self._get_office_xml_props(file_path))
            return meta
        except Exception as e:
            return {"error": str(e)}

    def _scrape_image(self, file_path):
        if not has_pil: return {"error": "Pillow not installed"}
        meta = {}
        try:
            img = Image.open(file_path)
            info = img._getexif()
            if info:
                for tag, value in info.items():
                    decoded = TAGS.get(tag, tag)
                    if decoded == "GPSInfo":
                        gps_data = {}
                        for t in value:
                            sub_decoded = GPSTAGS.get(t, t)
                            gps_data[sub_decoded] = value[t]
                        meta["GPS_Data"] = str(gps_data)
                    else:
                        meta[decoded] = str(value)
            meta["Format"] = img.format
            meta["Mode"] = img.mode
            meta["Size"] = f"{img.width}x{img.height}"
        except Exception as e:
            return {"error": str(e)}
        return meta

    def _extract_indicators(self, file_path):
        """Scrapes the binary/text content for URLs and Emails (Deep Content Scan)."""
        indicators = {"Emails": set(), "URLs": set()}
        try:
            with open(file_path, 'rb') as f:
                content = f.read().decode('utf-8', errors='ignore')
                emails = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', content)
                indicators["Emails"].update(emails)
                urls = re.findall(r'https?://[^\s<>"]+|www\.[^\s<>"]+', content)
                indicators["URLs"].update(urls)
        except: pass
        return {k: list(v) for k, v in indicators.items() if v}

    def _on_file_select(self, event):
        selection = self.file_list.curselection()
        if selection:
            file_path = self.file_list.get(selection[0])
            self.details_text.delete(1.0, tk.END)
            
            res_data = self.results.get(file_path, {})
            meta = res_data.get("metadata", {})
            indicators = res_data.get("indicators", {})

            if "error" in meta:
                self.details_text.insert(tk.END, f"❌ ERROR: {meta['error']}\n", "error")
                self.details_text.tag_config("error", foreground="red")
            else:
                self.details_text.insert(tk.END, f"📊 DEEP ANALYSIS: {os.path.basename(file_path)}\n")
                self.details_text.insert(tk.END, "="*60 + "\n\n")
                
                self.details_text.insert(tk.END, "[+] METADATA FIELDS:\n")
                for k, v in meta.items():
                    if v and v != 'None':
                        self.details_text.insert(tk.END, f"  {k.replace('/', ''):<25}: {v}\n")
                
                if indicators:
                    self.details_text.insert(tk.END, "\n" + "="*60 + "\n")
                    self.details_text.insert(tk.END, "[+] DISCOVERED INDICATORS (OSINT):\n")
                    for cat, items in indicators.items():
                        self.details_text.insert(tk.END, f"\n  {cat}:\n")
                        for item in items[:20]: # Show more indicators
                            self.details_text.insert(tk.END, f"    - {item}\n")
                        if len(items) > 20:
                            self.details_text.insert(tk.END, f"    ... and {len(items)-20} more.\n")

    def _start_analysis(self):
        path = self.path_var.get()
        if not path or not os.path.exists(path):
            messagebox.showerror("Error", "Please select a valid file or directory.")
            return

        self.analyze_btn.config(state=tk.DISABLED)
        self.status_var.set("Deep Analyzing...")
        self.file_list.delete(0, tk.END)
        self.results = {}
        
        threading.Thread(target=self._analyze_logic, args=(path,), daemon=True).start()

    def _analyze_logic(self, path):
        files_to_scan = []
        if os.path.isdir(path):
            for root, _, files in os.walk(path):
                for f in files:
                    if f.lower().endswith(('.pdf', '.docx', '.xlsx', '.pptx', '.jpg', '.jpeg', '.png')):
                        files_to_scan.append(os.path.join(root, f))
        else:
            files_to_scan.append(path)

        if not files_to_scan:
            self.root.after(0, lambda: messagebox.showinfo("Info", "No supported documents or images found."))
            self.root.after(0, lambda: self.analyze_btn.config(state=tk.NORMAL))
            self.root.after(0, lambda: self.status_var.set("Ready"))
            return

        for f_path in files_to_scan:
            ext = os.path.splitext(f_path)[1].lower()
            meta = {}
            if ext == '.pdf': meta = self._scrape_pdf(f_path)
            elif ext == '.docx': meta = self._scrape_docx(f_path)
            elif ext == '.xlsx': meta = self._scrape_xlsx(f_path)
            elif ext == '.pptx': meta = self._scrape_pptx(f_path)
            elif ext in ['.jpg', '.jpeg', '.png']: meta = self._scrape_image(f_path)
            
            if meta:
                indicators = self._extract_indicators(f_path)
                self.results[f_path] = {"metadata": meta, "indicators": indicators}
                self.root.after(0, lambda p=f_path: self.file_list.insert(tk.END, p))

        self.root.after(0, lambda: self.analyze_btn.config(state=tk.NORMAL))
        self.root.after(0, lambda: self.export_btn.config(state=tk.NORMAL))
        self.root.after(0, lambda: self.status_var.set(f"Done. Analyzed {len(self.results)} files."))

    def _export_results(self):
        if not self.results: return
        file_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON files", "*.json")])
        if file_path:
            with open(file_path, 'w') as f:
                json.dump(self.results, f, indent=4)
            messagebox.showinfo("Success", f"Results exported to {file_path}")

if __name__ == "__main__":
    root = tk.Tk()
    app = MetadataScraperGUI(root)
    root.mainloop()
